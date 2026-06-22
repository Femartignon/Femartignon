"""Corrige a validação corrompida (_xlfn._LONGTEXT, >255 chars) em DB_ATIVIDADES
que causava o aviso de recuperação do Excel. Reconstrói as 3 validações limpas."""
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation, DataValidationList

PATH = 'NEC_Eventos_Reestruturado_v4.xlsx'
wb = load_workbook(PATH)
at = wb['DB_ATIVIDADES']
last = int(at.tables['TblAtividades'].ref.split(':')[1][1:])

# Remover TODAS as validações da aba (incluindo a corrompida)
at.data_validations = DataValidationList()

# 1. Status (coluna H) — lista curta e válida
dv_status = DataValidation(type="list", formula1='"CONCLUÍDO,EM ANDAMENTO,PENDENTE,ATRASADO,-"',
    showDropDown=False, showInputMessage=True, showErrorMessage=True, allow_blank=True)
dv_status.promptTitle = "Status"
dv_status.sqref = f"H4:H{last}"
at.add_data_validation(dv_status)

# 2. Responsável (coluna E) — lista limpa, 57 chars (<255)
dv_resp = DataValidation(type="list",
    formula1='"Eventos,Fernando Azato,Janaina Silva,Marcio,Time de Marca"',
    showDropDown=False, showInputMessage=True, showErrorMessage=False, allow_blank=True)
dv_resp.promptTitle = "Responsável"
dv_resp.sqref = f"E4:E{last}"
at.add_data_validation(dv_resp)

# 3. Prazo (coluna F) — validação de data (recriada limpa)
dv_date = DataValidation(type="date", operator="greaterThanOrEqual", formula1="43831",
    allow_blank=True, showInputMessage=True, showErrorMessage=True)
dv_date.promptTitle = "Prazo da Atividade"
dv_date.prompt = "Insira a data no formato DD/MM/AAAA. Deixe em branco se não houver prazo."
dv_date.errorTitle = "Data inválida"
dv_date.error = "Insira uma data válida (DD/MM/AAAA) ou deixe em branco."
dv_date.sqref = f"F4:F{last}"
at.add_data_validation(dv_date)

wb.save(PATH)
print(f"✅ Validações reconstruídas em DB_ATIVIDADES (linhas 4..{last})")
print("   Status (H), Responsável (E), Prazo-data (F)")
